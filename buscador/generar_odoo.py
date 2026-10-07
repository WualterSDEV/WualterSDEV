import sys, re, unicodedata, openpyxl
from collections import Counter
from copy import copy
sys.path.insert(0, sys.argv[1])
from parse import parse; from norm import norm
D, ODOO, CRUCE, OUT = sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5]
def key(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper()
    return ' '.join(re.sub(r'[^A-Z0-9()/ ]', ' ', s).split())
b = parse(D+'/base.xls', 'base'); v = parse(D+'/venta.xls', 'venta')
for rr in b+v:
    for it in rr['items']:
        if it['ucosto'] == 'KILOS' and it['factor'] == 1000.0: it['ucosto'] = 'GRAMOS'
for r in b:
    if not r['nombre']: r['nombre'] = 'RB MASA HAMBURGUESA P'

ow = openpyxl.load_workbook(ODOO); ws = ow.active
hdr = [c for c in ws[1]]
print('tipos', Counter(r[2] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]))
ref = {}; used = set()        # nombre -> código Odoo
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0]: ref.setdefault(key(r[1]), r[0]); used.add(r[0])
    if r[5]: ref.setdefault(key(r[6]), r[5]); used.add(r[5])
loose = {}
for k, c in ref.items(): loose.setdefault((norm(k), k.startswith('RB '), k.startswith('(M')), set()).add(c)
def odoo_ref(name):
    k = key(name)
    if k in ref: return ref[k], 'exacto'
    s = loose.get((norm(k), k.startswith('RB '), k.startswith('(M')))
    if s and len(s) == 1: return next(iter(s)), 'parecido'
    return None, None

# selección: carta + toda la pastelería + recetas base necesarias
carta = set()
x = openpyxl.load_workbook(CRUCE, data_only=True)['Cruce Carta']
for row in x.iter_rows(min_row=2, values_only=True):
    if row[3]: carta.add(row[3])
vsel = [r for r in v if r['codigo'] in carta or r['area'] == 'PASTELERIA']
bidx = {}
for r in b: bidx.setdefault(norm(r['nombre']), r)
order = []; seen = set()
def visit(r):
    for it in r['items']:
        if it['tipo'] == 'Receta':
            s = bidx.get(norm(it['insumo']))
            if s and s['codigo'] not in seen and s is not r:
                seen.add(s['codigo']); visit(s); order.append(s)
for r in [r for r in b if r['area'] == 'PASTELERIA'] + vsel: visit(r)
for r in b:
    if r['area'] == 'PASTELERIA' and r['codigo'] not in seen: seen.add(r['codigo']); order.append(r)
recs = [(r, 'Base') for r in order] + [(r, 'Venta') for r in vsel]

out = openpyxl.Workbook(); out.remove(out.active)
def mk(title):
    sh = out.create_sheet(title)
    for j, h in enumerate(hdr, 1):
        c = sh.cell(1, j, h.value); c.font = copy(h.font); c.fill = copy(h.fill); c.alignment = copy(h.alignment); c.border = copy(h.border)
    for col, dim in ws.column_dimensions.items(): sh.column_dimensions[col].width = dim.width
    sh.freeze_panes = 'A2'; return sh
ok_sh, pend_sh = mk('Listo para subir'), mk('Pendientes')
pos = {ok_sh: 2, pend_sh: 2}
parecidos = []
from openpyxl.styles import PatternFill, Font
from openpyxl.comments import Comment
Y = PatternFill('solid', fgColor='FFEB9C'); R = PatternFill('solid', fgColor='FFC7CE')
rv = out.create_sheet('Revisar antes de subir')
for j, h in enumerate(['Qué', 'Nombre', 'Código sugerido', 'Usado en', 'Detalle'], 1): rv.cell(1, j, h).font = Font(bold=True)
issues = []
nuevos = {}
for r, t in recs:
    pid0, _ = odoo_ref(r['nombre'])
    if not pid0: nuevos[r['nombre']] = ''
ready = {}
for r, t in recs:
    pid0, _ = odoo_ref(r['nombre'])
    ready[r['nombre']] = bool(pid0) and all(odoo_ref(it['insumo'])[0] for it in r['items'])
for r, t in recs:
    o = ok_sh if ready[r['nombre']] else pend_sh; rown = pos[o]
    pid, how = odoo_ref(r['nombre'])
    if not pid:
        pid = None
        issues.append(('Producto (receta) no existe en Odoo', r['nombre'], '', '', f"Receta {t.lower()} Inforest {r['codigo']}. Crear el producto en Odoo, poner su código en la columna A de 'Pendientes' y subirla"))
    qty = 1.0 if t == 'Venta' else (r['factor'] or 1.0)
    first = True
    for it in r['items']:
        cref, chow = odoo_ref(it['insumo'])
        if not cref and it['insumo'] in nuevos: cref = nuevos[it['insumo']] or None; chow = 'nuevo'
        vals = ([pid, r['nombre'], 'Fabricar este producto', qty] if first else [None]*4) + \
               [it['ucosto'], cref, it['insumo'], round(it['cant'], 4) if isinstance(it['cant'], float) and it['cant'] != int(it['cant']) else int(it['cant'])]
        for j, val in enumerate(vals, 1): o.cell(rown, j, val)
        if first and how == 'parecido': parecidos.append((r['nombre'], pid))
        if chow == 'parecido': parecidos.append((it['insumo'], cref))
        if first and how != 'exacto' and how is not None:
            o.cell(rown, 2).comment = Comment('Nombre en Odoo un poco distinto; código tomado por similitud', 'Claude')
        if first and not how: o.cell(rown, 1).fill = Y
        if not cref:
            o.cell(rown, 6).fill = R
            issues.append(('Insumo sin código en Odoo', it['insumo'], '', r['nombre'], 'Crear o buscar el código en Odoo y escribirlo en la columna F'))
        elif chow == 'nuevo': o.cell(rown, 6).fill = Y
        elif chow == 'parecido': o.cell(rown, 6).comment = Comment('Código tomado por nombre parecido en Odoo', 'Claude')
        first = False; rown += 1
    if not r['items']:
        for j, val in enumerate([pid, r['nombre'], 'Fabricar este producto', qty], 1): o.cell(rown, j, val)
        rown += 1
    pos[o] = rown
# agrupar insumos sin código
agg = {}
for q, n, s, u, d in issues:
    k = (q, n)
    if k in agg: agg[k][3].add(u)
    else: agg[k] = [q, n, s, {u} if u else set(), d]
for i, (q, n, s, u, d) in enumerate(sorted(agg.values(), key=lambda a: (a[0], a[1])), 2):
    for j, val in enumerate([q, n, s, ', '.join(sorted(u)), d], 1): rv.cell(i, j, val)
for col, w in zip('ABCDE', [34, 40, 14, 60, 70]): rv.column_dimensions[col].width = w
out.move_sheet(rv, offset=0)
out.save(OUT)
print('recetas', len(recs), 'listas', sum(ready.values()), 'pendientes', len(recs)-sum(ready.values()))
print('parecidos', sorted(set(parecidos)))
print(Counter(q for q, *_ in agg.values()))
print('unidades', Counter(c.value for sh in (ok_sh, pend_sh) for c in sh['E'][1:]))
