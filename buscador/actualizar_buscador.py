"""Genera Buscador_Recetas.html a partir de los reportes de Inforest.

Uso:
  python actualizar_buscador.py recetas_base.xls recetas_venta.xls [recetas_odoo.xlsx] [cruce_carta.xlsx] [salida.html]

recetas_odoo.xlsx (opcional): exportación de listas de materiales de Odoo (como RECETAS_CUSCO.xlsx);
  de aquí salen los códigos AL... de productos e insumos para el Excel en formato Odoo.
cruce_carta.xlsx (opcional): Excel con la hoja 'Cruce Carta'; marca qué recetas están en la carta.
Requiere: pip install xlrd openpyxl
"""
import sys, os, re, json, unicodedata
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from parse import parse; from norm import norm
if len(sys.argv) < 3: sys.exit(__doc__)
arg = lambda i, d: sys.argv[i] if len(sys.argv) > i else os.path.join(HERE, '..', d)
b = parse(sys.argv[1], 'base'); v = parse(sys.argv[2], 'venta')
ODOO = arg(3, 'RECETAS_CUSCO_GRAMOS_ML.xlsx')
CRUCE = arg(4, 'RECETAS_INFOREST_CARTA_Y_POSTRES.xlsx')
OUT = arg(5, 'Buscador_Recetas.html')

def a_gramos(cant, unidad, factor=None):
    """Inforest a veces marca KILOS/LITROS con factor 1000: la cantidad ya está en gramos/ml."""
    if unidad in ('KILOS', 'LITROS'):
        nueva = 'GRAMOS' if unidad == 'KILOS' else 'MILILITROS'
        return (cant if factor == 1000.0 else (cant or 0) * 1000), nueva
    return cant, unidad

for rr in b + v:
    for it in rr['items']:
        it['cant'], it['ucosto'] = a_gramos(it['cant'], it['ucosto'], it['factor'])
for r in b:
    if not r['nombre']: r['nombre'] = 'RB MASA HAMBURGUESA P' if r['codigo'] == '00171' else '(sin nombre) ' + r['codigo']
    f = r['factor'] if isinstance(r['factor'], (int, float)) else None
    # Rinde = Factor expresado en la unidad de costo (ej. 2500 GRAMOS), no en la unidad de producción
    r['rinde'], r['urinde'] = a_gramos(f, r['ucosto'], f) if f is not None else (None, None)

carta = {}
if os.path.exists(CRUCE):
    import openpyxl
    x = openpyxl.load_workbook(CRUCE, data_only=True)['Cruce Carta']
    for row in x.iter_rows(min_row=2, values_only=True):
        if row[3]: carta.setdefault(row[3], (row[0], row[1], row[2]))

# Códigos de Odoo por nombre (productos y componentes)
def key(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper()
    return ' '.join(re.sub(r'[^A-Z0-9()/ ]', ' ', s).split())
ref, loose = {}, {}
if os.path.exists(ODOO):
    import openpyxl
    for r in openpyxl.load_workbook(ODOO, read_only=True).active.iter_rows(min_row=2, values_only=True):
        if r[0] and r[1]: ref.setdefault(key(r[1]), r[0])
        if r[5] and r[6]: ref.setdefault(key(r[6]), r[5])
    for k, c in ref.items(): loose.setdefault((norm(k), k.startswith('RB '), k.startswith('(M')), set()).add(c)
def odoo_ref(name):
    k = key(name)
    if k in ref: return ref[k]
    s = loose.get((norm(k), k.startswith('RB '), k.startswith('(M')))
    return next(iter(s)) if s and len(s) == 1 else None

bidx = {}
for r in b: bidx.setdefault(norm(r['nombre']), 'B' + r['codigo'])
def rnd(x): return round(x, 4) if isinstance(x, float) else x
out = []
for kind, recs in (('V', v), ('B', b)):
    for r in recs:
        items = []
        for it in r['items']:
            sub = bidx.get(norm(it['insumo'])) if it['tipo'] == 'Receta' else None
            if sub == kind + r['codigo']: sub = None
            items.append([it['insumo'], it['tipo'], rnd(it['cant']), it['ucosto'], rnd(it['precio']), rnd(it['subtotal']), sub, odoo_ref(it['insumo'])])
        c = carta.get(r['codigo']) if kind == 'V' else None
        out.append(dict(id=kind + r['codigo'], cod=r['codigo'], n=r['nombre'], t='Venta' if kind == 'V' else 'Base',
            cat=r['categoria'] or '', area=r['area'] or '', pv=rnd(r['venta']) if kind == 'V' and (r['venta'] or 0) >= 1 else None,
            costo=rnd(r['costo']), f=rnd(r.get('rinde')), up=r.get('urinde'), it=items, oid=odoo_ref(r['nombre']),
            carta=[c[0], c[1], c[2]] if c else None, postre=(r['area'] == 'PASTELERIA')))
data = json.dumps(out, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
lib = open(os.path.join(HERE, 'xlsx.full.min.js'), encoding='utf-8').read()
html = open(os.path.join(HERE, 'plantilla.html'), encoding='utf-8').read()
html = html.replace('/*XLSX*/', lib, 1).replace('/*DATA*/', data, 1).replace('/*FUENTE*/', 'Inforest 2026 · Cusco', 1)
open(OUT, 'w', encoding='utf-8').write(html)
print(f'{len(out)} recetas ({sum(1 for o in out if o["carta"])} en carta, {sum(1 for o in out if o["postre"])} de pastelería, '
      f'{sum(1 for o in out if o["oid"])} con código Odoo) -> {OUT}')
